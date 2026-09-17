#!/usr/bin/env python3
"""Differential C test harness: real ARM semantics (Unicorn) vs the actual
hardware (armv4t_2.circ, compiled to Verilog and run under iverilog).

For each .c file in a directory:
  1. Compile it against tools/harness_startup.S + tools/harness.ld -- the
     same toolchain/ABI as every other ROM in this repo, just with the
     halt-on-bx-to-0 convention (see harness_startup.S) so real, non-inlined
     function calls don't trip an early halt.
  2. Run the compiled image on Unicorn (a real, from-scratch ARM ISA
     implementation with no knowledge of this hardware's quirks) -- this is
     the oracle, not a hand-computed expected value.
  3. Run the SAME image on the actual hardware: `armv4t_2.circ` translated
     to Verilog by circ2v.py, executed under iverilog. This is not a model
     of the CPU, it IS the CPU's logic, mechanically compiled.
  4. Diff every word of the RAM window (0x1000-0x13FF, the only region a
     test is allowed to touch) between the two. Any difference is either a
     known gap (multiply, RRX, rotated-immediate carry, SWP, MSR/MRS --
     auto-detected from the disassembly) or a new, previously unknown bug.
  5. Write build/harness_report.xlsx: one row per test, full RAM diff on
     mismatch, a summary sheet, and a sheet listing every known-gap hit.

Usage:
    python3 tools/aggressive_c_harness.py c_tests/aggressive
    python3 tools/aggressive_c_harness.py c_tests/aggressive --circ armv4t_2.circ
"""
from __future__ import annotations

import argparse
import pathlib
import re
import struct
import subprocess
import sys
import tempfile

ROOT = pathlib.Path(__file__).resolve().parent.parent
COT_ROOT = pathlib.Path("/home/junaet/Documents/CustomCPU-COT")
sys.path.insert(0, str(COT_ROOT))

from tools import circ2v  # noqa: E402  (CustomCPU-COT's compiler; works generically)

CPU_FLAGS = ["-mcpu=arm7tdmi", "-marm"]
CFLAGS = CPU_FLAGS + ["-O2", "-ffreestanding", "-fno-builtin",
                      "-fomit-frame-pointer", "-fno-unwind-tables",
                      "-fno-asynchronous-unwind-tables"]
STARTUP = ROOT / "tools" / "harness_startup.S"
LDSCRIPT = ROOT / "tools" / "harness.ld"

RAM_BASE = 0x1000
RAM_WORDS = 256          # matches the hardware's real addrWidth=8 RAM
ROM_WORDS = 1024         # matches stage_IF's real addrWidth=10 ROM

KNOWN_GAPS = {
    "multiply":              r"\b(mul|mla|umull|umlal|smull|smlal)(s|eq|ne|cs|cc|mi|pl|vs|vc|hi|ls|ge|lt|gt|le)?\b",
    "rrx":                   r",\s*rrx\b",
    "swp":                   r"\bswpb?\b",
    "msr/mrs":               r"\b(msr|mrs)\b",
    "rotated-imm-carry (S-bit imm, needs manual check)": None,  # flagged separately
}


class CompileError(Exception):
    pass


def compile_c(src: pathlib.Path, workdir: pathlib.Path) -> tuple[list[int], str]:
    """Return (instruction words, objdump disassembly)."""
    obj = workdir / (src.stem + ".o")
    startup_o = workdir / "startup.o"
    elf = workdir / (src.stem + ".elf")
    binf = workdir / (src.stem + ".bin")

    def run(cmd):
        p = subprocess.run(cmd, capture_output=True, text=True)
        if p.returncode != 0:
            raise CompileError(p.stderr)

    run(["arm-none-eabi-gcc", *CFLAGS, "-c", str(src), "-o", str(obj)])
    if not startup_o.exists():
        run(["arm-none-eabi-gcc", *CPU_FLAGS, "-c", str(STARTUP), "-o", str(startup_o)])
    run(["arm-none-eabi-gcc", *CPU_FLAGS, "-nostdlib", "-Wl,--build-id=none",
         "-T", str(LDSCRIPT), str(startup_o), str(obj), "-o", str(elf)])
    run(["arm-none-eabi-objcopy", "-O", "binary", "--only-section=.text",
         str(elf), str(binf)])
    data = binf.read_bytes()
    if len(data) % 4:
        data += b"\x00" * (4 - len(data) % 4)
    words = list(struct.unpack("<%dI" % (len(data) // 4), data))
    if len(words) > ROM_WORDS:
        raise CompileError("%d words, ROM holds %d" % (len(words), ROM_WORDS))
    dump = subprocess.run(["arm-none-eabi-objdump", "-d", str(elf)],
                          capture_output=True, text=True).stdout
    return words, dump


def triage(disasm: str) -> list[str]:
    hits = []
    for name, pat in KNOWN_GAPS.items():
        if pat and re.search(pat, disasm, re.I):
            hits.append(name)
    # rotated-immediate carry: an S-suffixed data-processing/mov/mvn/cmp-class
    # instruction with a `#0x..` immediate whose top bits (an odd rotate) are
    # set -- heuristic, not exact; only flags the common case.
    for m in re.finditer(r"\b(movs|mvns|ands|orrs|eors|bics)\s+\w+,\s*(\w+,\s*)?#(0x[0-9a-f]+)", disasm, re.I):
        val = int(m.group(3), 16)
        if val > 0xFF and (val & (val - 1)) != 0:  # not a plain byte, not a clean power-of-2 pattern
            hits.append("rotated-imm-carry (heuristic)")
            break
    return hits


def run_unicorn(words: list[int], max_instrs: int = 2_000_000):
    from unicorn import Uc, UC_ARCH_ARM, UC_MODE_ARM, UC_HOOK_CODE
    from unicorn.arm_const import UC_ARM_REG_SP, UC_ARM_REG_PC

    mem_size = 0x2000  # 0x0000-0x1000 code, 0x1000-0x2000 data
    uc = Uc(UC_ARCH_ARM, UC_MODE_ARM)
    uc.mem_map(0, mem_size)
    blob = struct.pack("<%dI" % len(words), *words)
    uc.mem_write(0, blob)

    state = {"zero_hits": 0}

    def hook(uc, address, size, data):
        if address == 0:
            state["zero_hits"] += 1
            if state["zero_hits"] >= 2:
                uc.emu_stop()

    h = uc.hook_add(UC_HOOK_CODE, hook)
    try:
        uc.emu_start(0, mem_size, count=max_instrs)
    except Exception as e:                      # noqa: BLE001
        return {"halted": False, "error": str(e), "ram": {}}
    uc.hook_del(h)
    halted = state["zero_hits"] >= 2
    ram_bytes = uc.mem_read(RAM_BASE, RAM_WORDS * 4)
    ram = {i: w for i, w in enumerate(struct.unpack("<%dI" % RAM_WORDS, bytes(ram_bytes))) if w}
    return {"halted": halted, "error": None, "ram": ram}


_CPU_TOP_CACHE: dict[str, pathlib.Path] = {}


def build_cpu_top(circ_path: str) -> pathlib.Path:
    key = str(pathlib.Path(circ_path).resolve())
    if key in _CPU_TOP_CACHE:
        return _CPU_TOP_CACHE[key]
    outdir = pathlib.Path(tempfile.mkdtemp(prefix="hdl_"))
    v, decode_hex, prog_hex, meta = circ2v.build_verilog(circ_path)
    (outdir / "cpu_top.v").write_text(v + "\n")
    if decode_hex is not None:
        (outdir / "decode.hex").write_text("\n".join(decode_hex) + "\n")
    ram_arr = None
    for m in meta:
        f = m.split()
        if f[0] == "ram":
            ram_arr = f[1]
    (outdir / "meta.txt").write_text("\n".join(meta) + "\n")
    _CPU_TOP_CACHE[key] = outdir
    _CPU_TOP_CACHE[key + ".ram"] = ram_arr
    return outdir


def run_hardware(words: list[int], circ_path: str, max_cycles: int = 4000):
    outdir = build_cpu_top(circ_path)
    ram_arr = _CPU_TOP_CACHE[str(pathlib.Path(circ_path).resolve()) + ".ram"]
    depth = ROM_WORDS
    img = (list(words) + [0] * depth)[:depth]
    wd = pathlib.Path(tempfile.mkdtemp(prefix="hwrun_"))
    (wd / "prog.hex").write_text("\n".join("%08x" % (w & 0xffffffff) for w in img) + "\n")
    (wd / "decode.hex").write_bytes((outdir / "decode.hex").read_bytes())

    tb = """`timescale 1ns/1ps
module tb;
  reg CLK = 0, rst = 0;
  integer c;
  reg halted_flag = 0;
  wire halt;
  wire [31:0] o_pc_word_addr, o_instruction, o_alu_result, o_cpsr, o_bx_target, o_wb_wd;
  wire [0:0] o_alu_we, o_cond_pass, o_branch_taken, o_bl_taken, o_bx_taken, o_wb_we;
  wire [9:0] o_wa;
  cpu_top dut (
    .CLK(CLK), .rst(rst), .halt(halt),
    .o_pc_word_addr(o_pc_word_addr), .o_instruction(o_instruction),
    .o_alu_result(o_alu_result), .o_alu_we(o_alu_we), .o_cpsr(o_cpsr),
    .o_cond_pass(o_cond_pass), .o_branch_taken(o_branch_taken),
    .o_bl_taken(o_bl_taken), .o_bx_taken(o_bx_taken), .o_bx_target(o_bx_target),
    .o_wa(o_wa), .o_wb_wd(o_wb_wd), .o_wb_we(o_wb_we)
  );
  initial begin
    #1;
    for (c = 0; c < %d && !halted_flag; c = c + 1) begin
      CLK = 1; #1; CLK = 0; #1;
      if (o_bx_taken[0] === 1'b1 && o_bx_target === 0) halted_flag = 1;
    end
    if (halted_flag) $display("HW_HALTED");
    else $display("HW_TIMEOUT");
    for (c = 0; c < %d; c = c + 1)
      $display("RAM %%0d %%08h", c, dut.%s[c]);
    $finish;
  end
endmodule
""" % (max_cycles, RAM_WORDS, ram_arr)
    (wd / "tb.v").write_text(tb)
    subprocess.run(["iverilog", "-g2012", "-o", str(wd / "a.out"),
                    str(wd / "tb.v"), str(outdir / "cpu_top.v")],
                   check=True, capture_output=True)
    p = subprocess.run(["vvp", str(wd / "a.out")], cwd=wd,
                       capture_output=True, text=True, timeout=120)
    halted = "HW_HALTED" in p.stdout
    ram = {}
    for line in p.stdout.splitlines():
        if line.startswith("RAM "):
            _, k, v = line.split()
            if not re.fullmatch("x+", v, re.I) and int(v, 16) != 0:
                ram[int(k)] = int(v, 16)
    return {"halted": halted, "ram": ram}


def diff_ram(oracle: dict, hw: dict) -> list[tuple]:
    out = []
    for k in sorted(set(oracle) | set(hw)):
        a, b = oracle.get(k, 0), hw.get(k, 0)
        if a != b:
            out.append((k, a, b))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("srcdir")
    ap.add_argument("--circ", default=str(ROOT / "armv4t_2.circ"))
    ap.add_argument("--max-cycles", type=int, default=4000)
    ap.add_argument("--report", default=str(ROOT / "build" / "harness_report.xlsx"))
    args = ap.parse_args()

    files = sorted(pathlib.Path(args.srcdir).glob("*.c"))
    rows = []
    diff_rows = []
    print("compiling cpu_top.v from %s ..." % args.circ)
    build_cpu_top(args.circ)
    print("running %d C files ..." % len(files))

    for src in files:
        wd = pathlib.Path(tempfile.mkdtemp(prefix="ctest_"))
        row = {"test": src.stem, "file": str(src)}
        try:
            words, dump = compile_c(src, wd)
        except CompileError as e:
            row.update(status="COMPILE_ERROR", detail=str(e)[:2000])
            rows.append(row)
            print("[ERROR ] %-30s compile failed" % src.stem)
            continue

        row["words"] = len(words)
        row["gaps"] = ", ".join(triage(dump)) or ""

        oracle = run_unicorn(words)
        hw = run_hardware(words, args.circ, args.max_cycles)

        row["oracle_halted"] = oracle["halted"]
        row["hw_halted"] = hw["halted"]
        mismatches = diff_ram(oracle["ram"], hw["ram"])
        row["ram_words_touched"] = len(set(oracle["ram"]) | set(hw["ram"]))
        row["mismatch_count"] = len(mismatches)
        row["status"] = "PASS" if (not mismatches and oracle["halted"] == hw["halted"]) else "DIFFER"
        if not oracle["halted"]:
            row["status"] = "ORACLE_TIMEOUT"
        if not hw["halted"] and oracle["halted"]:
            row["status"] = "HW_TIMEOUT"

        for k, a, b in mismatches:
            diff_rows.append({"test": src.stem, "ram_word": "0x%02x" % k,
                              "bus_addr": "0x%04x" % (RAM_BASE + 4 * k),
                              "oracle": "0x%08x" % a, "hardware": "0x%08x" % b,
                              "known_gap": row["gaps"]})

        tag = row["status"]
        if row["gaps"] and tag == "DIFFER":
            tag = "DIFFER (known: %s)" % row["gaps"]
        print("[%-9s] %-30s %3dw  ram_diffs=%d  %s" %
             (row["status"], src.stem, row.get("words", 0), row["mismatch_count"], row["gaps"]))
        rows.append(row)

    write_report(rows, diff_rows, args.report)
    print("\nwrote %s" % args.report)
    passed = sum(1 for r in rows if r["status"] == "PASS")
    print("%d/%d PASS" % (passed, len(rows)))
    return 0 if passed == len(rows) else 1


def write_report(rows, diff_rows, path):
    sys.path.insert(0, str(pathlib.Path.home() / ".local/lib/python3.10/site-packages"))
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    def sheet(name, cols, data):
        ws = wb.new_sheet if False else wb.create_sheet(name)
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in data:
            ws.append([r.get(c, "") for c in cols])
        for i, c in enumerate(cols, 1):
            width = max(10, min(60, max([len(str(r.get(c, ""))) for r in data] + [len(c)]) + 2))
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        return ws

    del wb["Sheet"]
    summary_cols = ["test", "status", "words", "oracle_halted", "hw_halted",
                    "ram_words_touched", "mismatch_count", "gaps", "file"]
    ws = sheet("Summary", summary_cols, rows)
    fills = {"PASS": "C6EFCE", "DIFFER": "FFC7CE", "COMPILE_ERROR": "FFEB9C",
            "ORACLE_TIMEOUT": "FFEB9C", "HW_TIMEOUT": "FFC7CE"}
    status_col = summary_cols.index("status") + 1
    for i, r in enumerate(rows, 2):
        fill = fills.get(r.get("status"))
        if fill:
            ws.cell(i, status_col).fill = PatternFill("solid", fgColor=fill)

    sheet("RAM diffs", ["test", "ram_word", "bus_addr", "oracle", "hardware", "known_gap"], diff_rows)

    known = [r for r in rows if r.get("gaps")]
    sheet("Known-gap hits", summary_cols, known)

    errs = [r for r in rows if r["status"] == "COMPILE_ERROR"]
    sheet("Compile errors", ["test", "detail", "file"], errs)

    wb.save(path)


if __name__ == "__main__":
    sys.exit(main())
